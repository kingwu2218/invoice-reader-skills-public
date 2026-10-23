"""Score copied expectations against workbooks. Write score.json and accuracy-matrix.csv.

Correct empty cells are separate and excluded from headline accuracy denominators.
This is evidence on the named cases, never an accuracy estimate for user documents.
Usage: python eval/runner/score.py eval/runs/BATCH
"""
import argparse
import csv
import json
import hashlib
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
sys.path.insert(0, str(Path(__file__).resolve().parent))
import raw
import image_inputs

OUTCOMES = ('correct', 'correct_blank', 'miss', 'wrong_flagged', 'silent_error', 'invented')
DOCUMENT_SHEETS = {'Billing documents', 'Reference documents', 'Payments', 'Other', 'Documents'}


def rows(sheet):
    headers = [c.value for c in sheet[1]]
    return [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)
            if any(v not in (None, '') for v in row)]


def normalise(value, kind='text'):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if kind in ('amount', 'number'):
        try:
            return Decimal(str(value).replace(',', '')).normalize()
        except InvalidOperation:
            return ('text', str(value).strip())
    if kind == 'date':
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
    # Preserve case, punctuation and units: changing a printed identifier is an error.
    return ' '.join(str(value).split())


def flagged_labels(reasons):
    labels = set()
    for line in str(reasons or '').splitlines():
        if ':' in line:
            labels.update(part.strip() for part in line.split(':', 1)[0].split(','))
    return labels


def score_table(title, definition, actual, result):
    expected = definition['rows']
    specs = definition['columns']
    identity = definition['identity']
    unused = list(range(len(actual)))
    counts = Counter(expected=len(expected), actual=len(actual))
    for number, want in enumerate(expected):
        # Source/occurrence identity is independent of extracted field correctness.
        index = next((i for i in unused if all(normalise(actual[i].get(k)) == normalise(want.get(k))
                                               for k in identity)), None)
        got = actual[index] if index is not None else {}
        if index is not None:
            unused.remove(index)
            counts['matched'] += 1
        else:
            counts['missing'] += 1
        reasons = flagged_labels(got.get('Review reasons'))
        # Long custom observations carry the field name separately from their reason text.
        custom_flag = title == 'Requested values' and bool(got.get('Review reasons'))
        for spec in specs:
            label, kind = spec['label'], spec.get('kind', 'text')
            # A versioned reference may list several printed forms of one value, each
            # verified on the source (e.g. a name printed in two casings). The first is primary.
            wanted = want.get(label)
            accepted = [normalise(v, kind) for v in wanted] if isinstance(wanted, list) else [normalise(wanted, kind)]
            expected_value = accepted[0]
            actual_value = normalise(got.get(label), kind)
            if actual_value in accepted and None not in accepted:
                expected_value = actual_value
            if expected_value == actual_value:
                outcome = 'correct_blank' if expected_value is None else 'correct'
            elif actual_value is None:
                outcome = 'miss'
            elif expected_value is None:
                outcome = 'invented'
            elif label in reasons or any(line.startswith(label + ':') for line in str(got.get('Review reasons') or '').splitlines()) or (custom_flag and label in ('Value (as printed)', 'Unit')):
                outcome = 'wrong_flagged'
            else:
                outcome = 'silent_error'
            result['cells'][outcome] += 1
            cell_id = f'{title}/{number}/{label}'
            result['values'][cell_id] = str(actual_value)
            if outcome not in ('correct', 'correct_blank'):
                result['details'].append(dict(where=cell_id, outcome=outcome,
                                              expected=str(expected_value), got=str(actual_value)))
        if 'Needs review' in want:
            expected_flag, actual_flag = want['Needs review'] == 'Yes', got.get('Needs review') == 'Yes'
            label = ('true_positive' if actual_flag else 'false_negative') if expected_flag else (
                     'false_positive' if actual_flag else 'true_negative')
            result['flags'][label] += 1
        if title == 'Requested values' and got.get('Status') == 'printed':
            if not got.get('Printed evidence') or not got.get('Pages'):
                result['checks']['missing_evidence'] += 1
    counts['extra'] = len(unused)
    for i in unused:
        result['details'].append(dict(where=f'{title}/extra/{i}', outcome='extra_row', got=actual[i]))
        for spec in specs:
            if normalise(actual[i].get(spec['label']), spec.get('kind', 'text')) is not None:
                result['cells']['invented'] += 1
    result['tables'][title] = dict(counts)


def score_workbook(path, expected):
    result = dict(cells=Counter(), flags=Counter(), checks=Counter(), tables={}, values={}, details=[])
    with path.open('rb') as stream:
        book = load_workbook(stream)
    try:
        # Custom/line sheets store Document ID, not a duplicate Document type.
        # Join through actual document rows so same-page documents can be matched
        # independently of their extracted values and model-chosen ID spelling.
        document_metadata = defaultdict(list)
        document_rows, occurrences = {}, Counter()
        for title in book.sheetnames:
            if title not in DOCUMENT_SHEETS:
                continue
            document_rows[title] = rows(book[title])
            for row in document_rows[title]:
                identity = tuple(normalise(row.get(key)) for key in ('File', 'Pages', 'Document type'))
                occurrences[identity] += 1
                row['Document occurrence'] = str(occurrences[identity])
                if row.get('Document ID'):
                    document_metadata[(row.get('File'), row['Document ID'])].append(
                        {key: row.get(key) for key in ('Document type', 'Document occurrence', 'Pages')})
        for title, definition in expected['tables'].items():
            if title not in book.sheetnames:
                result['checks']['missing_sheets'] += 1
            actual = document_rows.get(title) if title in document_rows else (rows(book[title]) if title in book.sheetnames else [])
            if title in ('Requested data', 'Requested values', 'Line items', 'Statement lines'):
                for row in actual:
                    metadata = document_metadata.get((row.get('File'), row.get('Document ID')), [])
                    if len(metadata) == 1 and metadata[0]['Document type']:
                        # Line sheets lack Pages; use their unique parent document.
                        # Custom observation Pages is real evidence and must never
                        # be replaced by the parent document's broader page range.
                        row.update({key: value for key, value in metadata[0].items()
                                    if key != 'Pages' or title in ('Line items', 'Statement lines')})
            score_table(title, definition, actual, result)
            if definition.get('exact_columns') and title in book.sheetnames:
                headers = [c.value for c in book[title][1]]
                if headers != definition['exact_columns']:
                    result['checks']['wrong_columns'] += 1
        for title in DOCUMENT_SHEETS | {'Line items', 'Statement lines', 'Requested values', 'Requested data'}:
            if title not in expected['tables'] and title in book.sheetnames:
                extras = rows(book[title])
                if extras:
                    result['tables'][title] = dict(expected=0, actual=len(extras), matched=0, missing=0, extra=len(extras))
                    result['checks']['unexpected_rows'] += len(extras)
        if 'Input coverage' in book.sheetnames:
            result['checks']['unprocessed_inputs'] += sum(r.get('Status') != 'Processed' for r in rows(book['Input coverage']))
    finally:
        book.close()
    return result


def fraction(numerator, denominator):
    return {'numerator': numerator, 'denominator': denominator,
            'rate': numerator / denominator if denominator else None}


def metrics(result):
    cells = result['cells']
    correct = cells.get('correct', 0)
    wrong = cells.get('wrong_flagged', 0) + cells.get('silent_error', 0)
    supplied = correct + wrong + cells.get('invented', 0)
    printed = correct + wrong + cells.get('miss', 0)
    docs = Counter()
    for title, counts in result['tables'].items():
        if title in DOCUMENT_SHEETS:
            docs.update(counts)
    return {
        'field_accuracy': fraction(correct, printed + cells.get('invented', 0)),
        'field_precision': fraction(correct, supplied),
        'printed_value_recall': fraction(correct, printed),
        'document_recall': fraction(docs['matched'], docs['expected']),
        'document_precision': fraction(docs['matched'], docs['actual']),
    }


def score_batch(batch):
    results, matrix = [], []
    for run in sorted(batch.glob('*/run-*')):
        if not (run / 'meta.json').exists():
            continue
        meta = json.loads((run / 'meta.json').read_text(encoding='utf-8'))
        result = {'case': run.parent.name, 'run': run.name, 'valid': bool(meta.get('valid')), 'meta': meta}
        result['image_input_audit'] = image_inputs.audit(run, meta)
        try:
            if result['valid']:
                expected = json.loads((run / 'expected.json').read_text(encoding='utf-8'))
                result.update(score_workbook(run / 'output.xlsx', expected))
                result['metrics'] = metrics(result)
                result['raw_status'] = 'not_configured'
                if 'raw_tables' in expected:
                    try:
                        raw_bytes = (run/'raw-extraction.json').read_bytes()
                        if hashlib.sha256(raw_bytes).hexdigest() != meta.get('raw_extraction_hash'):
                            raise ValueError('Raw extraction hash is missing or does not match.')
                        result['raw'] = raw.score(json.loads(raw_bytes), expected['raw_tables'], score_table)
                        result['raw_metrics'] = metrics(result['raw'])
                        for name, title in (('line_item', 'Raw line items'), ('statement_line', 'Raw statement lines')):
                            counts = result['raw']['tables'].get(title, {})
                            result['raw_metrics'][name+'_recall'] = fraction(counts.get('matched', 0), counts.get('expected', 0))
                            result['raw_metrics'][name+'_precision'] = fraction(counts.get('matched', 0), counts.get('actual', 0))
                        result['raw_status'] = 'scored'
                    except (OSError, ValueError, KeyError, TypeError) as error:
                        result['raw_status'] = 'unavailable'
                        result['raw_error'] = str(error)
        except (OSError, ValueError, KeyError, BadZipFile) as error:
            result.update(valid=False, error=str(error))
        row = {k: meta.get(k, '') for k in ('family', 'tool', 'model_requested', 'model_reported', 'provider',
                                            'artifact_hash', 'dataset_hash', 'scenario', 'category', 'content_group')}
        row.update(case=result['case'], run=result['run'], valid=result['valid'])
        row['raw_status'] = result.get('raw_status', 'not_run')
        row['image_input_mode'] = meta.get('image_input_mode', 'not_captured')
        row['image_input_status'] = result['image_input_audit']['status']
        row['image_input_pages_numerator'] = result['image_input_audit']['numerator']
        row['image_input_pages_denominator'] = result['image_input_audit']['denominator']
        if result['valid']:
            row.update({outcome: result['cells'][outcome] for outcome in OUTCOMES})
            for name, value in result['metrics'].items():
                row[f'{name}_numerator'], row[f'{name}_denominator'], row[name] = value['numerator'], value['denominator'], value['rate']
            if result.get('raw_status') == 'scored':
                for outcome in OUTCOMES:
                    row['raw_'+outcome] = result['raw']['cells'][outcome]
                for name in ('field_accuracy', 'field_precision', 'printed_value_recall',
                             'line_item_recall', 'line_item_precision', 'statement_line_recall', 'statement_line_precision'):
                    value = result['raw_metrics'][name]
                    row['raw_'+name+'_numerator'] = value['numerator']
                    row['raw_'+name+'_denominator'] = value['denominator']
                    row['raw_'+name] = value['rate']
                row['raw_structural_failures'] = sum(c.get('missing',0)+c.get('extra',0) for c in result['raw']['tables'].values())
            for name in ('true_positive', 'true_negative', 'false_positive', 'false_negative'):
                row[f'review_{name}'] = result['flags'][name]
            row['structural_failures'] = sum(result['checks'].values()) + sum(
                c.get('missing', 0) + c.get('extra', 0) for c in result['tables'].values())
        matrix.append(row)
        results.append(result)
    if not results:
        raise ValueError('No runs found; an empty batch is not a successful evaluation.')
    by_case = defaultdict(list)
    for result in results:
        if result['valid']:
            by_case[result['case']].append(result)
    differing = 0
    for runs in by_case.values():
        keys = set().union(*(r['values'] for r in runs))
        differing += sum(len({r['values'].get(key, '<missing>') for r in runs}) > 1 for key in keys)
    report = {'runs': results, 'summary': {'runs': len(results), 'valid_runs': sum(r['valid'] for r in results),
                                          'cells_differing_between_runs': differing}}
    (batch / 'score.json').write_text(json.dumps(report, indent=2, default=str), encoding='utf-8')
    headers = list(dict.fromkeys(k for row in matrix for k in row))
    with (batch / 'accuracy-matrix.csv').open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=headers)
        writer.writeheader()
        writer.writerows(matrix)
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('batch', type=Path)
    args = parser.parse_args()
    report = score_batch(args.batch)
    print(json.dumps(report['summary']))
    failures = any(not r['valid'] or sum(r.get('checks', {}).values()) or
                   r['image_input_audit']['status'] == 'unavailable' or
                   r.get('raw_status') == 'unavailable' or
                   any(r.get('raw', {}).get('cells', {}).get(k, 0) for k in ('miss','wrong_flagged','silent_error','invented')) or
                   any(c.get('missing',0) or c.get('extra',0) for c in r.get('raw', {}).get('tables', {}).values()) or
                   r.get('flags', {}).get('false_positive', 0) or r.get('flags', {}).get('false_negative', 0) or
                   any(c.get('missing', 0) or c.get('extra', 0) for c in r.get('tables', {}).values()) or
                   any(r.get('cells', {}).get(k, 0) for k in ('miss', 'wrong_flagged', 'silent_error', 'invented'))
                   for r in report['runs'])
    return int(failures)


if __name__ == '__main__':
    raise SystemExit(main())
