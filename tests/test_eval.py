import copy
import importlib.util
import io
import json
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path

from openpyxl import load_workbook, Workbook
from test_custom_extraction import ROOT, coal_documents, process

spec = importlib.util.spec_from_file_location('scorer', ROOT / 'eval' / 'runner' / 'score.py')
score = importlib.util.module_from_spec(spec)
spec.loader.exec_module(score)


class Evaluation(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def workbook(self, scenario='coal-energy'):
        source, out = self.root / (scenario+'.json'), self.root / (scenario+'.xlsx')
        source.write_text(json.dumps(coal_documents(scenario)))
        folder = ROOT / 'eval' / 'scenarios' / scenario
        with redirect_stdout(io.StringIO()):
            process.main(['--extraction', str(source), '--out', str(out), '--custom', str(folder/'custom.md')])
        return out, json.loads((folder/'expected.json').read_text())

    def test_listed_source_forms_are_accepted_but_other_text_stays_strict(self):
        definition = {'identity': ['File'], 'columns': [{'label': 'Payer'}],
                      'rows': [{'File': 'a.pdf', 'Payer': ['ACME CO. LIMITED', 'Acme Co. Limited']}]}
        for got, outcome in (('ACME CO. LIMITED', 'correct'), ('Acme Co. Limited', 'correct'),
                             ('ACME CO LIMITED', 'silent_error'), ('acme co. limited', 'silent_error'), (None, 'miss')):
            with self.subTest(got=got):
                result = dict(cells=__import__('collections').Counter(), flags=__import__('collections').Counter(),
                              checks=__import__('collections').Counter(), tables={}, values={}, details=[])
                score.score_table('Payments', definition, [{'File': 'a.pdf', 'Payer': got}], result)
                self.assertEqual(dict(result['cells']), {outcome: 1})
        plain = {'identity': ['File'], 'columns': [{'label': 'Payer'}], 'rows': [{'File': 'a.pdf', 'Payer': 'ACME CO. LIMITED'}]}
        result = dict(cells=__import__('collections').Counter(), flags=__import__('collections').Counter(),
                      checks=__import__('collections').Counter(), tables={}, values={}, details=[])
        score.score_table('Payments', plain, [{'File': 'a.pdf', 'Payer': 'Acme Co. Limited'}], result)
        self.assertEqual(dict(result['cells']), {'silent_error': 1})

    def test_requested_rows_join_document_type_without_matching_values(self):
        book=Workbook(); book.remove(book.active)
        docs=book.create_sheet('Billing documents')
        docs.append(['Document ID','File','Pages','Document type'])
        docs.append(['model-a','source.pdf','1','Invoice'])
        docs.append(['model-b','source.pdf','1','Receipt'])
        values=book.create_sheet('Requested values')
        values.append(['Document ID','File','Document pages','Field key','Occurrence','Status','Value (as printed)'])
        values.append(['model-b','source.pdf','1','number','1','absent',None])
        values.append(['model-a','source.pdf','1','number','1','unclear',None])
        path=self.root/'joined.xlsx'; book.save(path); book.close()
        expected={'tables':{'Requested values':{
            'identity':['File','Document pages','Document type','Field key','Occurrence'],
            'columns':[{'label':'Status'}],
            'rows':[{'File':'source.pdf','Document pages':'1','Document type':t,'Field key':'number','Occurrence':'1','Status':s}
                    for t,s in [('Invoice','unclear'),('Receipt','absent')]]}}}
        result=score.score_workbook(path,expected)
        self.assertEqual(result['cells']['correct'],2)
        self.assertEqual(result['tables']['Requested values']['matched'],2)
        # Ambiguous duplicate IDs must not silently choose one document type.
        book=load_workbook(path); book['Billing documents'].append(['model-a','source.pdf','1','Receipt'])
        book.save(path); book.close()
        result=score.score_workbook(path,expected)
        self.assertEqual(result['tables']['Requested values']['missing'],1)

    def test_both_personalised_scenarios(self):
        for scenario in ('coal-payment', 'coal-energy'):
            out, expected = self.workbook(scenario)
            result = score.score_workbook(out, expected)
            self.assertEqual(result['details'], [], scenario)
            self.assertFalse(any(result['checks'].values()), scenario)
            self.assertEqual(score.metrics(result)['field_accuracy']['rate'], 1)

    def test_extra_invoice_fails(self):
        out, expected = self.workbook()
        book = load_workbook(out)
        sheet = book['Billing documents']
        sheet.append([c.value for c in sheet[2]])
        book.save(out)
        book.close()
        result = score.score_workbook(out, expected)
        self.assertEqual(result['tables']['Billing documents']['extra'], 1)
        self.assertGreater(result['cells']['invented'], 0)
        self.assertLess(score.metrics(result)['document_precision']['rate'], 1)

    def test_missing_document_and_custom_observation_fail(self):
        out, expected = self.workbook()
        book = load_workbook(out)
        book['Other'].delete_rows(2)
        book['Requested values'].delete_rows(5)
        book.save(out)
        book.close()
        result = score.score_workbook(out, expected)
        self.assertEqual(result['tables']['Other']['missing'], 1)
        self.assertEqual(result['tables']['Requested values']['missing'], 1)

    def test_correct_blanks_do_not_inflate_accuracy(self):
        out, expected = self.workbook()
        result = score.score_workbook(out, expected)
        self.assertGreater(result['cells']['correct_blank'], 0)
        metric = score.metrics(result)['field_accuracy']
        self.assertEqual(metric['denominator'], result['cells']['correct'])
        self.assertLess(metric['denominator'], sum(result['cells'].values()))

    def test_wrong_unit_is_silent_error(self):
        out, expected = self.workbook()
        book = load_workbook(out)
        sheet = book['Requested values']
        labels = [c.value for c in sheet[1]]
        sheet.cell(4, labels.index('Unit')+1).value = 'MJ/kg'
        book.save(out)
        book.close()
        result = score.score_workbook(out, expected)
        self.assertGreater(result['cells']['silent_error'], 0)

    def test_matrix_preserves_invalid_runs_and_denominators(self):
        out, expected = self.workbook()
        batch = self.root / 'batch'
        run = batch / 'coal-energy' / 'run-1'
        run.mkdir(parents=True)
        (run/'output.xlsx').write_bytes(out.read_bytes())
        (run/'expected.json').write_text(json.dumps(expected))
        (run/'meta.json').write_text(json.dumps({'valid': True, 'model_requested': 'test-only'}))
        bad = batch / 'coal-energy' / 'run-2'
        bad.mkdir()
        (bad/'meta.json').write_text('{"valid":false}')
        report = score.score_batch(batch)
        self.assertEqual(report['summary']['valid_runs'], 1)
        self.assertEqual(report['summary']['runs'], 2)
        self.assertIn('field_accuracy_denominator', (batch/'accuracy-matrix.csv').read_text(encoding='utf-8-sig'))

    def test_empty_batch_is_not_success(self):
        with self.assertRaisesRegex(ValueError, 'No runs'):
            score.score_batch(self.root)

    def test_case_and_punctuation_are_not_silently_normalised(self):
        self.assertNotEqual(score.normalise('AB-01'), score.normalise('ab01'))
