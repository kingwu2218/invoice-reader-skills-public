"""Frozen public replies, exact-session continuity and preservation of earlier results."""
import io
import importlib.util
import json
from contextlib import redirect_stdout
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'eval/runner'))
import clarification
import conversation
import run as baseline
import score
import process
from openpyxl import Workbook

CASE = 'dash-origin-printed-v1'


class ClarificationEvaluation(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name); self.base = self.root/'conversation'
        self.base.mkdir(); self.work = self.root/'work'; self.work.mkdir()
        self.scenario = self.root/'eval/scenarios'/CASE
        shutil.copytree(ROOT/'eval/scenarios'/CASE, self.scenario)
        (self.root/'eval/fixtures').mkdir()
        shutil.copyfile(ROOT/'eval/fixtures'/f'{CASE}.pdf', self.root/'eval/fixtures'/f'{CASE}.pdf')
        shutil.copyfile(self.root/'eval/fixtures'/f'{CASE}.pdf', self.work/f'{CASE}.pdf')
        shutil.copyfile(self.scenario/'custom.md', self.work/'custom.md')
        self.package = self.work/'.agents/skills/invoice-reader-openai'; self.package.mkdir(parents=True)
        (self.package/'SKILL.md').write_text('test package')
        (self.work/'output.xlsx').write_bytes(b'original workbook')
        self.repo = patch.object(baseline, 'REPO', self.root); self.repo.start(); self.addCleanup(self.repo.stop)
        self.state = dict(schema_version=1, phase='response_received', session_id='exact-session',
                          workspace=str(self.work), case=CASE, case_id=CASE, turns=[],
                          artifact_hash=baseline.skillsets.tree_hash(self.package),
                          clarification_lock=clarification.lock_case(CASE))
        conversation.save(self.base/'state.json', self.state)

    def test_scripted_reply_keeps_answers_out_of_workspace_and_uses_exact_session(self):
        def invoke(state, prompt, folder, session):
            self.assertEqual(session, 'exact-session')
            self.assertIn('page 1 only', prompt)
            self.assertNotIn('negative 720', prompt)
            self.assertFalse(list(self.work.glob('expected*')))
            self.assertEqual((self.work/'output.xlsx').read_bytes(), b'original workbook')
            self.assertEqual(state['output_name'], 'output-first.xlsx')
            return dict(completed=True, agent_messages=['Page 2 still needs confirmation.'], session_ids=[session]), {}, conversation.snapshot(self.work)
        with patch.object(conversation,'invoke',side_effect=invoke), patch.object(conversation,'finish') as finish, redirect_stdout(io.StringIO()):
            clarification.followup(self.base, 'first')
            finish.assert_called_once()
        state = conversation.load_state(self.base)
        turn = state['turns'][-1]
        self.assertEqual(turn['actor'], 'scripted_public_fixture')
        self.assertTrue(turn['audit']['prior_workbooks_unchanged'])
        self.assertEqual(state['prior_output_hashes']['output.xlsx'], clarification.digest(self.work/'output.xlsx'))

    def test_out_of_order_repeated_or_unlocked_phase_cannot_invoke(self):
        with patch.object(conversation,'invoke') as invoke:
            with self.assertRaises(ValueError): clarification.followup(self.base, 'second')
            self.state['clarification_lock'] = None
            conversation.save(self.base/'state.json', self.state)
            with self.assertRaises(ValueError): clarification.followup(self.base, 'first')
            invoke.assert_not_called()

    def test_changed_future_reference_or_package_cannot_invoke(self):
        for target in (self.scenario/'expected-second.json', self.package/'SKILL.md'):
            with self.subTest(target=target.name), patch.object(conversation,'invoke') as invoke:
                original = target.read_bytes(); target.write_bytes(original+b' ')
                with self.assertRaises(ValueError): clarification.followup(self.base, 'first')
                invoke.assert_not_called(); target.write_bytes(original)

    def test_unsafe_or_duplicate_output_names_are_rejected(self):
        path = self.scenario/'clarification.json'; plan = json.loads(path.read_text())
        for output in ('../output.xlsx', 'output.xlsx', 'output-first.xlsx'):
            plan['phases'][1]['output'] = output
            path.write_text(json.dumps(plan))
            with self.assertRaises(ValueError): clarification.lock_case(CASE)

    def test_unexpected_previous_workbook_is_not_accepted_as_history(self):
        (self.work/'output (2).xlsx').write_bytes(b'extra')
        with patch.object(conversation,'invoke') as invoke:
            with self.assertRaises(ValueError): clarification.followup(self.base, 'first')
            invoke.assert_not_called()

    def test_finish_preserves_named_output_and_flags_changed_earlier_workbook(self):
        # Exercise real artifact capture with a minimal manifest and named result.
        run = self.base/'extractions'/CASE/'run-1'; run.mkdir(parents=True)
        (run/'expected.json').write_text('{}')
        self.state.update(output_name='output-first.xlsx', prior_output_hashes={'output.xlsx':clarification.digest(self.work/'output.xlsx')})
        (self.work/'output-first.xlsx').write_bytes(b'new workbook')
        source = self.work/f'{CASE}.pdf'
        import hashlib
        self.state['dataset_hash'] = hashlib.sha256(source.read_bytes()+(self.work/'custom.md').read_text().encode()+b'{}').hexdigest()
        (self.work/'output.xlsx').write_bytes(b'changed')
        info = dict(completed=True, exit_code=0, models_reported_by_cli=[], elapsed_seconds=1)
        conversation.finish(self.state, self.base, info)
        meta = json.loads((run/'meta.json').read_text())
        self.assertFalse(meta['valid'])
        self.assertIn('earlier workbook changed', meta['output_error'])
        self.assertEqual((run/'output.xlsx').read_bytes(), b'new workbook')

    def test_line_pages_join_unique_parent_without_overwriting_observation_pages(self):
        book = Workbook(); docs = book.active; docs.title = 'Billing documents'
        docs.append(['File','Pages','Document ID','Document type'])
        docs.append(['pack.pdf','1','first','Invoice']); docs.append(['pack.pdf','2-3','second','Invoice'])
        lines = book.create_sheet('Line items'); lines.append(['File','Document ID','Line','Line total'])
        lines.append(['pack.pdf','second',1,720]); lines.append(['pack.pdf','first',1,360])
        obs = book.create_sheet('Requested values'); obs.append(['File','Document ID','Pages','Value (as printed)'])
        obs.append(['pack.pdf','second','3','720-'])
        expected = {'tables': {
            'Line items': {'identity':['File','Pages','Document type','Document occurrence','Line'],
                           'columns':[{'label':'Line total','kind':'amount'}], 'rows':[
                               {'File':'pack.pdf','Pages':page,'Document type':'Invoice','Document occurrence':'1','Line':'1','Line total':amount}
                               for page, amount in [('1','360'),('2-3','720')]]},
            'Requested values': {'identity':['File','Pages'], 'columns':[{'label':'Value (as printed)','kind':'text'}],
                                 'rows':[{'File':'pack.pdf','Pages':'3','Value (as printed)':'720-'}]}}}
        path = self.root/'join.xlsx'; book.save(path)
        self.assertFalse(score.score_workbook(path, expected)['details'])
        lines['D2'], lines['D3'] = 360, 720; book.save(path)
        self.assertEqual(score.score_workbook(path, expected)['cells']['silent_error'], 2)
        docs['C3'] = 'first'; book.save(path); book.close()
        self.assertEqual(score.score_workbook(path, expected)['tables']['Line items']['missing'], 2)

    def test_reference_v2_covers_all_required_sheets_and_frozen_v1_is_unchanged(self):
        spec = importlib.util.spec_from_file_location('dash_origin_fixture', ROOT/'eval/tools/make_dash_origin_case.py')
        fixture = importlib.util.module_from_spec(spec); spec.loader.exec_module(fixture)
        for phase in ('pending', 'first', 'second'):
            with self.subTest(phase=phase):
                name = 'expected.json' if phase == 'pending' else f'expected-{phase}.json'
                self.assertEqual(fixture.expected(phase), json.loads((ROOT/'eval/scenarios'/CASE/name).read_text()))
                expected = json.loads((ROOT/'eval/scenarios/dash-origin-printed-v2'/name).read_text())
                self.assertEqual(fixture.expected_v2(phase), expected)
                docs = []
                for page, amount in [('1','360'),('2','720')]:
                    notation = {'origin':'computer-generated'}
                    if page == '1' and phase != 'pending': notation.update(confirmed_as='end-mark', confirmation='Fixture: page 1 positive')
                    if page == '2' and phase == 'second': notation.update(confirmed_as='negative', confirmation='Fixture: page 2 negative')
                    docs.append({'document_id':page, 'file':fixture.FILE, 'pages':page, 'document_type':'Invoice',
                                 'fields':{'document_title':'INVOICE','document_number':'SYN-ORIGIN-0'+page,'total':amount+'-','currency':'HKD'},
                                 'amount_notation':{'total':notation},
                                 'line_items':[{'quantity':'1','unit_price':amount+'.00','line_total':amount+'.00'}],
                                 'custom_fields':{'printed_total':[{'status':'printed','value':amount+'-','unit':'HKD','pages':page,'evidence':'Total (HKD) '+amount+'-'}]}})
                source = self.root/(phase+'.json'); source.write_text(json.dumps({'documents':docs}))
                output = self.root/(phase+'.xlsx')
                with redirect_stdout(io.StringIO()):
                    self.assertEqual(process.main(['--extraction',str(source),'--out',str(output),'--custom',str(self.scenario/'custom.md')]), 0)
                result = score.score_workbook(output, expected)
                self.assertFalse(result['details']); self.assertFalse(any(result['checks'].values()))


if __name__ == '__main__':
    unittest.main()
