"""Keep raw evidence, quantities and unclear values outside numeric overrides."""
from copy import deepcopy
from contextlib import redirect_stdout
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'core/scripts'))
sys.path.insert(0,str(ROOT/'eval/runner'))
import process
import score
from openpyxl import load_workbook

CASE=ROOT/'eval/processor-cases/amount-dash-v1'


class AmountDashPipeline(unittest.TestCase):
    def test_missing_origin_stays_text_without_changing_raw_or_custom(self):
        source = (CASE/'extraction.json').read_bytes()
        with tempfile.TemporaryDirectory() as scratch:
            folder = Path(scratch); inp = folder/'input.json'; inp.write_bytes(source)
            out = folder/'output.xlsx'
            with redirect_stdout(io.StringIO()):
                self.assertEqual(process.main(['--extraction',str(inp),'--out',str(out),
                                               '--custom',str(CASE/'custom.md')]), 0)
            expected = json.loads((CASE/'expected-unknown.json').read_text())
            result = score.score_workbook(out, expected)
            self.assertFalse(result['details']); self.assertFalse(any(result['checks'].values()))
            self.assertEqual(inp.read_bytes(), source)
            book = load_workbook(out)
            info = dict(book['Run info'].values)
            self.assertEqual(info['Trailing amount dash convention'], 'auto')
            self.assertEqual(info['Trailing amount dash basis'], 'per-field origin and recorded user replies')
            self.assertIn('pending user confirmation', score.rows(book['Billing documents'])[0]['Review reasons'])
            self.assertNotIn('confirmed', score.rows(book['Billing documents'])[0]['Review reasons'])
            book.close()

    def test_frozen_matrix_and_raw_custom_values_survive(self):
        source=(CASE/'extraction.json').read_bytes()
        for mode in ('unknown','end-mark','negative'):
            with self.subTest(mode=mode), tempfile.TemporaryDirectory() as scratch:
                folder=Path(scratch); inp=folder/'input.json'; inp.write_bytes(source)
                out=folder/'output.xlsx'
                with redirect_stdout(io.StringIO()):
                    code=process.main(['--extraction',str(inp),'--out',str(out),'--custom',str(CASE/'custom.md'),
                                       '--amount-dash',mode])
                self.assertEqual(code,0); self.assertEqual(inp.read_bytes(),source)
                expected=json.loads((CASE/f'expected-{mode}.json').read_text())
                result=score.score_workbook(out,expected)
                self.assertEqual(result['details'],[])
                self.assertFalse(any(result['checks'].values()))
                self.assertEqual(result['flags']['true_positive'],3)
                book=load_workbook(out)
                info=dict(book['Run info'].values)
                self.assertEqual(info['Trailing amount dash convention'],mode)
                self.assertEqual(info['Trailing amount dash basis'],'explicit batch option')
                invoice=score.rows(book['Billing documents'])[0]
                if mode=='unknown':
                    self.assertIn('negative amount or an end mark',invoice['Review reasons'])
                else:
                    self.assertIn("converted '125-'",invoice['Review reasons'])
                book.close()

    def test_convention_cannot_override_unclear_or_quantity(self):
        data=json.loads((CASE/'extraction.json').read_text())
        data['documents'][0]['unclear']=['total']
        data['documents'][0]['line_items'][0]['quantity']='1-'
        original=deepcopy(data)
        for mode in ('end-mark','negative'):
            result=process.Processor(None,process._load_toml('review-reasons.toml'),process.load_columns(),
                                     process.default_mapping(),process.custom_fields.read_fields(CASE/'custom.md'),
                                     amount_dash=mode).run(data)
            self.assertIsNone(result.rows['Billing documents'][0]['total'])
            self.assertEqual(result.line_items[0]['quantity'],'1-')
            self.assertEqual(result.line_items[0]['line_check'],'Not checkable')
            self.assertEqual(data,original)

    def process_data(self, data, **kwargs):
        process.extraction_io.validate(data)
        return process.Processor(None, process._load_toml('review-reasons.toml'), process.load_columns(),
                                 process.default_mapping(), process.custom_fields.read_fields(CASE/'custom.md'),
                                 **kwargs).run(data)

    def mixed_data(self):
        data = json.loads((CASE/'extraction.json').read_text())
        invoice, cheque, statement = data['documents']
        invoice['handwritten'] = True
        invoice['amount_notation'] = {'total': {'origin': 'computer-generated'}}
        invoice['line_items'][0]['amount_notation'] = {
            'unit_price': {'origin': 'handwritten'}, 'line_total': {'origin': 'computer-generated'}}
        cheque['amount_notation'] = {'amount': {'origin': 'handwritten'}}
        statement['statement_lines'][0]['amount_notation'] = {
            'debit': {'origin': 'handwritten'}, 'balance': {'origin': 'unknown'}}
        return data

    def test_mixed_origins_are_scoped_to_fields_in_every_table(self):
        data = self.mixed_data(); original = deepcopy(data)
        result = self.process_data(data)
        invoice = result.rows['Billing documents'][0]
        self.assertEqual(invoice['total'], '125-')
        self.assertEqual(invoice['arithmetic_check'], 'Not checkable')
        self.assertEqual(result.line_items[0]['unit_price'], 125)
        self.assertEqual(result.line_items[0]['line_total'], '125-')
        self.assertEqual(result.line_items[0]['line_check'], 'Not checkable')
        self.assertEqual(result.rows['Payments'][0]['amount'], 125)
        self.assertIn('handwritten notation identified for this field', result.rows['Payments'][0]['review_reasons'])
        self.assertEqual(result.statement_lines[0]['debit'], 125)
        self.assertEqual(result.statement_lines[0]['balance'], '125-')
        self.assertEqual(data, original)

    def test_document_handwritten_flag_never_substitutes_for_field_origin(self):
        data = json.loads((CASE/'extraction.json').read_text())
        for doc in data['documents']:
            doc['handwritten'] = True
        result = self.process_data(data)
        self.assertEqual(result.rows['Billing documents'][0]['total'], '125-')
        self.assertEqual(result.rows['Payments'][0]['amount'], '125-')
        self.assertEqual(result.line_items[0]['line_total'], '125-')
        self.assertEqual(result.statement_lines[0]['balance'], '125-')

    def test_scoped_user_reply_resolves_printed_field_without_changing_handwriting(self):
        for convention, expected in [('end-mark', 125), ('negative', -125)]:
            with self.subTest(convention=convention):
                data = self.mixed_data()
                data['documents'][0]['amount_notation']['total'].update(
                    confirmed_as=convention, confirmation='Fixture user confirms only invoice Total: '+convention)
                original = deepcopy(data)
                result = self.process_data(data)
                self.assertEqual(result.rows['Billing documents'][0]['total'], expected)
                self.assertIn('recorded user reply:', result.rows['Billing documents'][0]['review_reasons'])
                self.assertEqual(result.line_items[0]['unit_price'], 125)
                self.assertEqual(result.line_items[0]['line_total'], '125-')
                self.assertEqual(result.rows['Payments'][0]['amount'], 125)
                self.assertEqual(data, original)
                held = self.process_data(data, amount_dash='unknown')
                self.assertEqual(held.rows['Billing documents'][0]['total'], '125-')

    def test_invalid_origin_metadata_and_conflicting_batch_fail(self):
        invalid = [[], {'total': 'handwritten'}, {'total': {'origin': 'maybe'}},
                   {'total': {'origin': 'handwritten', 'extra': True}},
                   {'total': {'origin': 'computer-generated', 'confirmed_as': 'negative'}},
                   {'total': {'origin': 'unknown', 'confirmation': 'No selected sign'}},
                   {'missing': {'origin': 'handwritten'}},
                   {'document_title': {'origin': 'handwritten'}}]
        for metadata in invalid:
            with self.subTest(metadata=metadata):
                data = self.mixed_data(); data['documents'][0]['amount_notation'] = metadata
                with self.assertRaises(process.ExtractionError):
                    self.process_data(data)
        data = self.mixed_data()
        data['documents'][0]['line_items'][0]['amount_notation'] = {'quantity': {'origin': 'handwritten'}}
        with self.assertRaises(process.ExtractionError):
            self.process_data(data)
        data = self.mixed_data()
        data['documents'][0]['amount_notation']['total'].update(confirmed_as='negative', confirmation='Fixture reply')
        with self.assertRaises(process.ExtractionError):
            self.process_data(data, amount_dash='end-mark')

    def test_unclear_stays_blank_and_raw_grading_does_not_treat_metadata_as_values(self):
        import raw
        data = self.mixed_data()
        data['documents'][1]['unclear'] = ['amount']
        self.assertIsNone(self.process_data(data).rows['Payments'][0]['amount'])
        tables = raw.tables(data)
        self.assertNotIn('item.amount_notation', tables['Raw line items'][0])
        self.assertEqual(tables['Raw line items'][0]['item.unit_price'], '125-')
        self.assertEqual(tables['Raw documents'][1]['fields.amount'], '125-')
