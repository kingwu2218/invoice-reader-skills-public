import copy
import io
import json
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'core' / 'scripts'))
import buckets
import custom_fields
import extraction
import process
from checks import check_document, NOT_CHECKABLE
from decimal import Decimal


def coal_documents(scenario='coal-energy'):
    invoice = {'file': 'coal-pack.pdf', 'pages': '1', 'document_type': 'Invoice',
               'fields': {'document_title': 'INVOICE', 'document_number': 'COAL-001',
                          'seller_name': 'Example Fuel Lab Ltd', 'buyer_name': 'Tallowmere Consulting Ltd',
                          'subtotal': '100.00', 'tax_amount': '0.00', 'total': '105.00'},
               'line_items': [{'quantity': '1', 'unit_price': '100.00', 'line_total': '100.00'}]}
    certificate = {'file': 'coal-pack.pdf', 'pages': '2', 'document_type': 'Lab certificate',
                   'fields': {'document_title': 'LAB CERTIFICATE'}, 'description': 'Coal certificate'}
    if scenario == 'coal-energy':
        invoice['custom_fields'] = {'coal_energy': [{'status': 'absent'}], 'moisture': [{'status': 'absent'}]}
        certificate['custom_fields'] = {'coal_energy': [
            {'status': 'printed', 'value': '6,200', 'unit': 'kcal/kg', 'context': 'Sample C-01', 'pages': '2',
             'evidence': 'Sample C-01; Energy Level of Coal: 6,200 kcal/kg'},
            {'status': 'printed', 'value': '6,100', 'unit': 'kcal/kg', 'context': 'Sample C-02', 'pages': '2',
             'evidence': 'Sample C-02; Energy Level of Coal: 6,100 kcal/kg'}], 'moisture': [{'status': 'absent'}]}
    else:
        invoice['custom_fields'] = {
            'invoice_ref': [{'status': 'printed', 'value': 'COAL-001', 'pages': '1', 'evidence': 'Invoice No.: COAL-001'}],
            'payment_due': [{'status': 'printed', 'value': '15 Oct 2026', 'pages': '1', 'evidence': 'Payment due date: 15 Oct 2026'}]}
        certificate['custom_fields'] = {'invoice_ref': [{'status': 'absent'}], 'payment_due': [{'status': 'absent'}]}
    return {'schema_version': 1, 'documents': [invoice, certificate]}


class CustomExtraction(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.custom = ROOT / 'eval' / 'scenarios' / 'coal-energy' / 'custom.md'
        self.specs = custom_fields.read_fields(self.custom)

    def run_process(self, payload, scenario='coal-energy', manifest=False):
        source = self.root / 'records.json'
        source.write_text(json.dumps(payload), encoding='utf-8')
        out = self.root / (scenario + '.xlsx')
        args = ['--extraction', str(source)]
        if manifest:
            p = self.root / 'manifest.json'
            p.write_text(json.dumps({'version': 1, 'inputs': [{'file': 'coal-pack.pdf', 'extraction': source.name}]}))
            args = ['--manifest', str(p)]
        with redirect_stdout(io.StringIO()):
            code = process.main(args + ['--out', str(out), '--custom', str(ROOT/'eval'/'scenarios'/scenario/'custom.md')])
        return code, out

    def test_arbitrary_measurements_and_repeated_units(self):
        code, out = self.run_process(coal_documents(), manifest=True)
        self.assertEqual(code, 0)
        book = load_workbook(out)
        self.addCleanup(book.close)
        self.assertEqual(book.sheetnames[0], 'Requested data')
        headers = [c.value for c in book['Requested data'][1]]
        value = book['Requested data'].cell(3, headers.index('Energy Level of Coal')+1).value
        self.assertEqual(value, 'Sample C-01: 6,200 kcal/kg\nSample C-02: 6,100 kcal/kg')
        detail = list(book['Requested values'].values)
        self.assertEqual(len(detail)-1, 5)
        self.assertIn('Printed evidence', detail[0])
        self.assertIn('Input coverage', book.sheetnames)

    def test_same_source_different_requested_columns(self):
        _, a = self.run_process(coal_documents())
        _, b = self.run_process(coal_documents('coal-payment'), 'coal-payment')
        books = [load_workbook(p) for p in (a, b)]
        for book in books:
            self.addCleanup(book.close)
        self.assertNotEqual(list(books[0]['Requested data'].values)[0], list(books[1]['Requested data'].values)[0])

    def test_unclear_cannot_contain_a_value(self):
        doc = coal_documents()['documents'][1]
        doc['custom_fields']['coal_energy'] = [{'status': 'unclear', 'value': '6200', 'pages': '2'}]
        with self.assertRaisesRegex(buckets.CustomError, 'stay blank'):
            custom_fields.observations(doc, self.specs)

    def test_missing_evidence_refused(self):
        doc = coal_documents()['documents'][1]
        del doc['custom_fields']['coal_energy'][0]['evidence']
        with self.assertRaisesRegex(buckets.CustomError, 'evidence'):
            custom_fields.observations(doc, self.specs)

    def test_evidence_cannot_point_outside_document(self):
        doc = coal_documents()['documents'][1]
        doc['custom_fields']['coal_energy'][0]['pages'] = '1'
        with self.assertRaisesRegex(buckets.CustomError, 'belong'):
            custom_fields.observations(doc, self.specs)

    def test_page_ranges(self):
        self.assertEqual(extraction.page_numbers('1, 3-4'), {1, 3, 4})
        for pages in ('0', '4-2', '100001', 'page two'):
            with self.assertRaises(extraction.ExtractionError):
                extraction.page_numbers(pages)

    def test_missing_and_unrequested_keys_refused(self):
        for change in ('missing', 'extra'):
            doc = coal_documents()['documents'][1]
            if change == 'missing':
                del doc['custom_fields']['moisture']
            else:
                doc['custom_fields']['invented_key'] = [{'status': 'absent'}]
            with self.assertRaises(buckets.CustomError):
                custom_fields.observations(doc, self.specs)

    def test_one_does_not_drop_extra_observations(self):
        fields = copy.deepcopy(self.specs)
        fields[0]['values'] = 'one'
        with self.assertRaises(buckets.CustomError):
            custom_fields.observations(coal_documents()['documents'][1], fields)

    def test_formula_like_custom_value_is_text(self):
        payload = coal_documents()
        payload['documents'][1]['custom_fields']['coal_energy'][0]['value'] = '=HYPERLINK("bad")'
        _, out = self.run_process(payload)
        book = load_workbook(out)
        self.addCleanup(book.close)
        found = [c for row in book['Requested values'] for c in row if c.value == '=HYPERLINK("bad")']
        self.assertEqual(len(found), 1)
        self.assertEqual(found[0].data_type, 's')

    def test_malformed_field_table_refused(self):
        p = self.root / 'custom.md'
        for row in ('| bad-key | Test | Find this | one | no |', '| x | Test | Find this | all | no |'):
            p.write_text(buckets.default_text().replace('|---|---|---|---|---|', '|---|---|---|---|---|\n' + row))
            with self.assertRaises(buckets.CustomError):
                custom_fields.read_fields(p)

    def test_bare_document_and_wrong_types_refused(self):
        for part in ({'file': 'x.pdf'}, {'documents': []}, {'documents': 'bad'}, {'documents': [{'file': 'x.pdf', 'document_type': 'Invoice', 'fields': {'total': 10}}]}):
            with self.assertRaises(extraction.ExtractionError):
                extraction.validate(part)

    def test_manifest_ignores_stale_json(self):
        (self.root / 'stale.json').write_text('{not even valid json')
        code, _ = self.run_process(coal_documents(), manifest=True)
        self.assertEqual(code, 0)

    def test_manifest_partial_output_is_explicit(self):
        p = self.root / 'manifest.json'
        p.write_text(json.dumps({'version': 1, 'inputs': [{'file': 'broken.pdf', 'extraction': None, 'reason': 'Cannot read PDF'}]}))
        out = self.root / 'partial.xlsx'
        with redirect_stdout(io.StringIO()):
            code = process.main(['--manifest', str(p), '--out', str(out)])
        self.assertEqual(code, 3)
        self.assertTrue(out.exists())

    def test_manifest_rejects_escape_and_source_mismatch(self):
        for item in ({'file': 'a.pdf', 'extraction': '../outside.json'}, {'file': 'a.pdf', 'extraction': 'records.json'}):
            (self.root / 'records.json').write_text(json.dumps(coal_documents()))
            p = self.root / 'manifest.json'
            p.write_text(json.dumps({'version': 1, 'inputs': [item]}))
            with self.assertRaises(extraction.ExtractionError):
                extraction.read_manifest(p)

    def test_duplicate_documents_refused(self):
        payload = coal_documents()
        payload['documents'].append(copy.deepcopy(payload['documents'][0]))
        p = self.root / 'records.json'
        p.write_text(json.dumps(payload))
        with self.assertRaisesRegex(extraction.ExtractionError, 'Duplicate'):
            extraction.read(p)

    def test_unknown_tax_cannot_pass_as_zero(self):
        result, _ = check_document({'subtotal': Decimal('100'), 'total': Decimal('100')}, [], blocked={'tax_amount'})
        self.assertEqual(result, NOT_CHECKABLE)

    def test_unclear_total_is_not_replaced(self):
        doc = coal_documents()['documents'][0]
        doc.pop('custom_fields')
        doc['unclear'] = ['total']
        doc['fields']['balance_due'] = '80'
        source = self.root / 'one.json'
        source.write_text(json.dumps({'documents': [doc]}))
        out = self.root / 'one.xlsx'
        with redirect_stdout(io.StringIO()):
            self.assertEqual(process.main(['--extraction', str(source), '--out', str(out)]), 0)
        book = load_workbook(out)
        self.addCleanup(book.close)
        rows = list(book['Billing documents'].values)
        self.assertIsNone(rows[1][rows[0].index('Total')])
