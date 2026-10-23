"""Tests for the deterministic skill scripts. Run: python -m unittest discover tests"""

import json
import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parents[1] / "core" / "scripts"
sys.path.insert(0, str(SCRIPTS))

from openpyxl import load_workbook  # noqa: E402

import buckets  # noqa: E402
import new_profile  # noqa: E402
import process  # noqa: E402
from amounts import parse_amount  # noqa: E402
from checks import FAILED, NOT_CHECKABLE, PASSED, check_document, check_line  # noqa: E402
from company_profile import PURCHASE, SALES, UNCLEAR, direction, read_profile  # noqa: E402
from dates import parse_date  # noqa: E402

D = Decimal


def sheet_rows_raw(sheet):
    return [tuple(c.value for c in r) for r in sheet.iter_rows()]


def sheet_rows(sheet):
    headers = [c.value for c in sheet[1]]
    return [dict(zip(headers, (c.value for c in r))) for r in sheet.iter_rows(min_row=2)]


class Amounts(unittest.TestCase):
    def test_terminal_dash_requires_field_origin_or_explicit_convention(self):
        for printed in ('125-', '1,250.50-', 'HK$ 125 -', '.50-'):
            with self.subTest(printed=printed):
                self.assertEqual(parse_amount(printed, trailing_dash='unknown'), (None, 'ambiguous'))
                positive, status = parse_amount(printed, trailing_dash='end-mark')
                negative, negative_status = parse_amount(printed, trailing_dash='negative')
                self.assertEqual((status, negative_status), ('ok', 'ok'))
                self.assertEqual(positive, -negative)
                self.assertGreater(positive, 0)
                self.assertEqual(parse_amount(printed), (None, 'ambiguous'))
                self.assertEqual(parse_amount(printed, origin='computer-generated'), (None, 'ambiguous'))
                self.assertEqual(parse_amount(printed, origin='handwritten'), (positive, 'ok'))
        self.assertEqual(parse_amount('125-', trailing_dash='end-mark'), (D('125'), 'ok'))
        self.assertEqual(parse_amount('125-', trailing_dash='negative'), (D('-125'), 'ok'))
        self.assertEqual(parse_amount('125-', trailing_dash=None), (None, 'unreadable'))

    def test_terminal_dash_does_not_repair_other_notation(self):
        for convention in ('auto', 'unknown', 'end-mark', 'negative'):
            for printed in ('-125-', '+125-', '(125-)', '125- CR', '125--', '12,50-',
                            '125/2-', '125/-', '12?5-', '125–', '125—', '125%-'):
                with self.subTest(convention=convention, printed=printed):
                    self.assertEqual(parse_amount(printed, trailing_dash=convention), (None, 'unreadable'))
        with self.assertRaises(ValueError):
            parse_amount('125-', trailing_dash='guess')

    def test_printed_forms(self):
        cases = {
            "1,234.50": D("1234.50"), "$1,234.50": D("1234.50"), "USD 99.00": D("99.00"),
            "(1,234.50)": D("-1234.50"), "1,234.50 CR": D("-1234.50"), "-125.00": D("-125.00"),
            "0.00": D("0.00"), "HK$ 12.30": D("12.30"), "£5": D("5"),
        }
        for printed, expected in cases.items():
            with self.subTest(printed=printed):
                self.assertEqual(parse_amount(printed), (expected, "ok"))

    def test_blank_stays_blank(self):
        self.assertEqual(parse_amount(""), (None, "blank"))
        self.assertEqual(parse_amount(None), (None, "blank"))

    def test_cheque_terminal_stroke_is_context_scoped(self):
        self.assertEqual(parse_amount('8,750/', cheque_amount=True), (D('8750'), 'ok'))
        self.assertEqual(parse_amount('8,750/'), (None, 'unreadable'))
        for value in ('3/4', '8,750//', '8,750/2', '8?750/', '12,34/'):
            self.assertEqual(parse_amount(value, cheque_amount=True), (None, 'unreadable'))

    def test_unreadable_is_not_guessed(self):
        for printed in ("1.234,50", "12,34", "1,2345.00", "about 50", "1?3.00"):
            with self.subTest(printed=printed):
                self.assertEqual(parse_amount(printed), (None, "unreadable"))


class Dates(unittest.TestCase):
    def test_unambiguous(self):
        for printed in ("6 Oct 2026", "6th October, 2026", "October 6, 2026",
                        "2026-10-06", "06-Oct-26", "Oct. 6 2026"):
            with self.subTest(printed=printed):
                self.assertEqual(parse_date(printed), (date(2026, 10, 6), "ok"))

    def test_numeric_needs_date_order(self):
        self.assertEqual(parse_date("03/04/2026"), (None, "ambiguous"))
        self.assertEqual(parse_date("03/04/2026", "Day/Month/Year"), (date(2026, 4, 3), "ok"))
        self.assertEqual(parse_date("03/04/2026", "Month/Day/Year"), (date(2026, 3, 4), "ok"))
        self.assertEqual(parse_date("25/10/2026"), (date(2026, 10, 25), "ok"))
        self.assertEqual(parse_date("04/04/2026"), (date(2026, 4, 4), "ok"))

    def test_impossible(self):
        self.assertEqual(parse_date("31 Feb 2026"), (None, "unreadable"))
        self.assertEqual(parse_date("next Tuesday"), (None, "unreadable"))

    def test_spaced_date_keeps_ambiguity_and_calendar_checks(self):
        self.assertEqual(parse_date('24 11 2026'), (date(2026, 11, 24), 'ok'))
        self.assertEqual(parse_date('03 04 2026'), (None, 'ambiguous'))
        self.assertEqual(parse_date('03 04 2026', 'Day/Month/Year'), (date(2026, 4, 3), 'ok'))
        self.assertEqual(parse_date('31 02 2026'), (None, 'unreadable'))


class Checks(unittest.TestCase):
    def test_line(self):
        self.assertEqual(check_line(D("3"), D("10.00"), None, D("30.00")), PASSED)
        self.assertEqual(check_line(D("3"), D("10.00"), D("5.00"), D("25.00")), PASSED)
        self.assertEqual(check_line(D("3"), D("10.00"), None, D("31.00")), FAILED)
        self.assertEqual(check_line(None, D("10.00"), None, D("30.00")), NOT_CHECKABLE)

    def test_document(self):
        amounts = {"subtotal": D("100"), "tax_amount": D("5"), "total": D("105")}
        self.assertEqual(check_document(amounts, [D("60"), D("40")]), (PASSED, []))
        self.assertEqual(check_document(amounts, [D("60"), D("41")]), (FAILED, ["subtotal_mismatch"]))
        self.assertEqual(check_document({**amounts, "total": D("106")}, []), (FAILED, ["total_mismatch"]))

    def test_missing_input_is_not_checkable_not_passed(self):
        self.assertEqual(check_document({"total": D("105")}, []), (NOT_CHECKABLE, []))
        self.assertEqual(check_document({"subtotal": D("100")}, [None]), (NOT_CHECKABLE, []))

    def test_no_subtotal_uses_lines(self):
        amounts = {"tax_amount": D("5"), "total": D("105")}
        self.assertEqual(check_document(amounts, [D("60"), D("40")]), (PASSED, []))


def _profile(folder, companies, date_order="Not sure", required=()):
    path = Path(folder) / "company-profile.xlsx"
    new_profile.build(path)
    book = load_workbook(path)
    for name, aliases in companies:
        book["Company"].append([name, aliases])
    book["Settings"]["B2"] = date_order
    for row in book["Columns"].iter_rows(min_row=2):
        if row[0].value in required:
            row[2].value = "Yes"
    book.save(path)
    return path


class Profile(unittest.TestCase):
    def test_direction(self):
        with tempfile.TemporaryDirectory() as folder:
            profile = read_profile(_profile(folder, [("ABC Trading Limited", "ABC Trading Co")]))
        self.assertEqual(direction(profile, "ABC Trading Ltd.", "Other Co"), (SALES, None))
        self.assertEqual(direction(profile, "Other Co", "abc trading co"), (PURCHASE, None))
        self.assertEqual(direction(profile, "Other Co", "Someone"), (UNCLEAR, "direction_none"))
        self.assertEqual(direction(profile, "ABC Trading Ltd", "ABC Trading Co"),
                         (UNCLEAR, "direction_both"))
        self.assertEqual(direction(None, "A", "B"), (UNCLEAR, "no_profile"))


class EndToEnd(unittest.TestCase):
    def run_process(self, extraction, profile_rows=(("Example Buyer Ltd", ""),),
                    sheet="Billing documents", custom=None, **profile_args):
        folder = tempfile.mkdtemp()
        profile = _profile(folder, profile_rows, **profile_args)
        source = Path(folder) / "extraction.json"
        source.write_text(json.dumps(extraction), encoding="utf-8")
        out = Path(folder) / "out.xlsx"
        args = ["--extraction", str(source), "--out", str(out),
                "--profile", str(profile), "--model-reported", "test-model"]
        if custom is not None:
            (Path(folder) / "custom.md").write_text(custom, encoding="utf-8")
            args += ["--custom", str(Path(folder) / "custom.md")]
        self.assertEqual(process.main(args), 0)
        book = load_workbook(out)
        return book, sheet_rows(book[sheet])

    def invoice(self, **fields):
        base = {"document_title": "INVOICE", "document_number": "INV-1", "issue_date": "6 Oct 2026",
                "seller_name": "Example Seller Ltd", "buyer_name": "Example Buyer Ltd",
                "subtotal": "100.00", "tax_amount": "5.00", "total": "105.00"}
        base.update(fields)
        return {"documents": [{
            "file": "inv.pdf", "pages": "1", "document_type": "Invoice",
            "type_evidence": "Title reads INVOICE", "fields": base,
            "line_items": [{"description": "Widget", "quantity": "2", "unit_price": "50.00",
                            "line_total": "100.00"}]}]}

    def test_clean_invoice(self):
        book, rows = self.run_process(self.invoice())
        row = rows[0]
        self.assertEqual(row["Direction"], PURCHASE)
        self.assertEqual(row["Total"], 105)
        self.assertEqual(row["Issue date"].date(), date(2026, 10, 6))
        self.assertEqual(row["Arithmetic check"], PASSED)
        self.assertEqual(row["Needs review"], "No")
        self.assertIsNone(row["Due date"])

    def test_formula_text_stays_text(self):
        book, rows = self.run_process(self.invoice(seller_name="=HYPERLINK(\"x\")", reference="@SUM(A1)"))
        cell = next(c for c in book["Billing documents"][2] if c.value == "=HYPERLINK(\"x\")")
        self.assertEqual(cell.data_type, "s")
        self.assertTrue(cell.quotePrefix)

    def test_negative_amount_stays_a_number(self):
        extraction = self.invoice(subtotal="-100.00", tax_amount="-5.00", total="-105.00")
        extraction["documents"][0]["document_type"] = "Credit note"
        extraction["documents"][0]["line_items"] = []
        _book, rows = self.run_process(extraction)
        self.assertEqual(rows[0]["Total"], -105)

    def test_mismatch_and_unclear_are_flagged(self):
        extraction = self.invoice(total="106.00", due_date="7 Nov 2026")
        extraction["documents"][0]["unclear"] = ["due_date"]
        _book, rows = self.run_process(extraction)
        row = rows[0]
        self.assertEqual(row["Arithmetic check"], FAILED)
        self.assertIsNone(row["Due date"], "an unclear value must be left blank")
        self.assertIn("Subtotal, Tax amount, Total: the total doesn't equal", row["Review reasons"])
        self.assertIn("Due date: unclear on the page", row["Review reasons"])
        self.assertEqual(row["Needs review"], "Yes")

    def test_ambiguous_date_and_required_field(self):
        _book, rows = self.run_process(self.invoice(issue_date="03/04/2026"), required=("Due date",))
        reasons = rows[0]["Review reasons"]
        self.assertIn("Issue date: this date could be read two ways", reasons)
        self.assertIn("Due date: your company profile requires this", reasons)
        self.assertEqual(rows[0]["Issue date"], "03/04/2026")

    def test_missing_title_is_flagged(self):
        extraction = self.invoice()
        del extraction["documents"][0]["fields"]["document_title"]
        _book, rows = self.run_process(extraction)
        self.assertIn("Document title: no document title is printed", rows[0]["Review reasons"])

    def test_balance_due_checked_only_when_no_total(self):
        extraction = self.invoice(balance_due="105.00")
        del extraction["documents"][0]["fields"]["total"]
        _book, rows = self.run_process(extraction)
        self.assertEqual(rows[0]["Arithmetic check"], PASSED)
        self.assertEqual(rows[0]["Total"], 105, "with no printed total, the balance due stands for it")

        extraction = self.invoice(balance_due="150.00")
        del extraction["documents"][0]["fields"]["total"]
        _book, rows = self.run_process(extraction)
        self.assertEqual(rows[0]["Total"], 150)
        self.assertIn("Subtotal, Tax amount, Total, Balance due: no total is printed",
                      rows[0]["Review reasons"])

        extraction = self.invoice(balance_due="80.00")
        _book, rows = self.run_process(extraction)
        self.assertEqual(rows[0]["Balance due"], 80, "a balance after payments is kept as printed")
        self.assertEqual(rows[0]["Arithmetic check"], PASSED, "Total is the one checked")

    def test_invoice_number(self):
        _book, rows = self.run_process(self.invoice())
        self.assertEqual(rows[0]["Invoice No."], "INV-1", "an invoice's own number fills Invoice No.")
        extraction = self.invoice(document_title="CREDIT NOTE", document_number="CN-7",
                                  invoice_number="INV-1")
        extraction["documents"][0]["document_type"] = "Credit note"
        _book, rows = self.run_process(extraction)
        self.assertEqual((rows[0]["Document number"], rows[0]["Invoice No."]), ("CN-7", "INV-1"))
        extraction["documents"][0]["fields"].pop("invoice_number")
        _book, rows = self.run_process(extraction)
        self.assertIsNone(rows[0]["Invoice No."], "a credit note's own number is not an invoice number")

    def test_receipt_is_extracted(self):
        extraction = self.invoice(document_title="OFFICIAL RECEIPT")
        extraction["documents"][0]["document_type"] = "Receipt"
        book, rows = self.run_process(extraction, sheet="Reference documents")
        self.assertEqual(rows[0]["Document type"], "Receipt")
        self.assertEqual(rows[0]["Document title"], "OFFICIAL RECEIPT")
        self.assertEqual(sheet_rows(book["Billing documents"]), [])

    def test_custom_moves_receipts_to_billing(self):
        extraction = self.invoice(document_title="RECEIPT")
        extraction["documents"][0]["document_type"] = "Receipt"
        custom = buckets.default_text(confirmed=True).replace(
            "| Receipt | Reference documents |", "| Receipt | Billing documents |")
        book, rows = self.run_process(extraction, custom=custom)
        self.assertEqual(rows[0]["Document type"], "Receipt")
        info = dict(sheet_rows_raw(book["Run info"]))
        self.assertEqual(info["Document buckets"], "custom.md (confirmed)")

    def test_payment_goes_to_payments(self):
        extraction = {"documents": [{
            "file": "chq.pdf", "pages": "1", "document_type": "Cheque",
            "type_evidence": "Pay ... or order; cheque number line",
            "fields": {"payment_date": "3 Oct 2026", "bank_name": "Example Bank",
                       "cheque_number": "000123", "payee": "Example Buyer Ltd",
                       "payer": "Someone Else Ltd", "amount": "250.00",
                       "amount_in_words": "Two hundred and fifty pounds only"}}]}
        _book, rows = self.run_process(extraction, sheet="Payments")
        self.assertEqual(rows[0]["Direction"], process.MONEY_IN)
        self.assertEqual(rows[0]["Amount"], 250)
        self.assertIn("Document title: no document title is printed", rows[0]["Review reasons"])

    def test_unknown_type_goes_to_other_and_is_flagged(self):
        extraction = {"documents": [{
            "file": "x.pdf", "pages": "1", "document_type": "Tenancy agreement",
            "type_evidence": "Title reads TENANCY AGREEMENT", "description": "A lease for an office",
            "fields": {"document_title": "TENANCY AGREEMENT", "document_date": "1 Sep 2026",
                       "parties": "Example Buyer Ltd; Landlord Ltd"}}]}
        _book, rows = self.run_process(extraction, sheet="Other")
        self.assertEqual(rows[0]["Needs review"], "Yes")
        self.assertIn("Not a document type the reader handles in detail", rows[0]["Review reasons"])

    def test_statement_lines(self):
        extraction = self.invoice(document_title="STATEMENT", balance_due="300.00")
        doc = extraction["documents"][0]
        doc["document_type"] = "Statement of account"
        doc["line_items"] = []
        doc["statement_lines"] = [{"date": "1 Sep 2026", "reference": "INV-1", "debit": "300.00",
                                   "balance": "300.00"}]
        book, rows = self.run_process(extraction, sheet="Reference documents")
        lines = sheet_rows(book["Statement lines"])
        self.assertEqual(lines[0]["Debit"], 300)
        self.assertEqual(lines[0]["Document number"], "INV-1")


    def test_unknown_field_is_refused(self):
        folder = tempfile.mkdtemp()
        source = Path(folder) / "extraction.json"
        bad = self.invoice(invoice_no="1")
        bad["set_aside"] = []
        source.write_text(json.dumps(bad), encoding="utf-8")
        code = process.main(["--extraction", str(source), "--out", str(Path(folder) / "o.xlsx")])
        self.assertEqual(code, 2)


class Custom(unittest.TestCase):
    def write(self, text):
        path = Path(tempfile.mkdtemp()) / "custom.md"
        path.write_text(text, encoding="utf-8")
        return path

    def test_default_is_unconfirmed_and_complete(self):
        mapping, confirmed = buckets.read_custom(self.write(buckets.default_text()))
        self.assertFalse(confirmed)
        self.assertEqual(mapping, buckets.default_mapping())

    def test_shape_rules(self):
        base = buckets.default_text(confirmed=True)
        bad = base.replace("| Cheque | Payments |", "| Cheque | Billing documents |")
        with self.assertRaisesRegex(buckets.CustomError, "Cheque can't go on Billing documents"):
            buckets.read_custom(self.write(bad))
        unsure = base.replace("| Invoice | Billing documents |",
                              "| Invoice | Billing documents |\n| Unsure | Reference documents |")
        with self.assertRaisesRegex(buckets.CustomError, "unsure"):
            buckets.read_custom(self.write(unsure))
        added = base.replace("| Invoice | Billing documents |",
                             "| Invoice | Billing documents |\n| Tenancy agreement | Other |")
        mapping, _ = buckets.read_custom(self.write(added))
        self.assertEqual(buckets.place("tenancy agreement", mapping), ("Other", "other"))
        self.assertEqual(buckets.place("Unsure", mapping), ("Other", "other"))


if __name__ == "__main__":
    unittest.main()
