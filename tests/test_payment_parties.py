"""Party omissions/swaps and same-page invoices must not hide behind amount accuracy."""
import copy
from contextlib import closing, redirect_stdout
import importlib.util
import io
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
from test_raw_eval import ROOT, payload, raw, score, process


class PaymentParties(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.work = Path(self.temp.name)
        self.case = ROOT/'eval/scenarios/format-parties-v2'
        self.expected = json.loads((self.case/'expected.json').read_text(encoding='utf-8'))
        self.data = payload()
        self.data['documents'][1]['fields'].update(
            payee='Example Workshop Ltd', payer='Tallowmere Consulting Ltd', bank_name='Example Bank')
        self.data['documents'][1]['custom_fields']['payment_amount'][0]['unit'] = 'HKD'

    def workbook(self, data, name):
        source = self.work/(name+'.json'); source.write_text(json.dumps(data), encoding='utf-8')
        output = self.work/(name+'.xlsx')
        with redirect_stdout(io.StringIO()):
            self.assertEqual(process.main(['--extraction',str(source),'--out',str(output),
                                           '--custom',str(self.case/'custom.md')]), 0)
        return output

    def test_new_reference_covers_parties_without_changing_original(self):
        spec = importlib.util.spec_from_file_location('party_case', ROOT/'eval/tools/make_party_case.py')
        module = importlib.util.module_from_spec(spec); spec.loader.exec_module(module)
        self.assertEqual(module.expected(2), self.expected)
        self.assertEqual(module.expected(1), json.loads((ROOT/'eval/scenarios/format-parties-v1/expected.json').read_text()))
        for version in (1, 2):
            with self.assertRaises(FileExistsError): module.main(version)
        raw_result = raw.score(self.data, self.expected['raw_tables'], score.score_table)
        book_result = score.score_workbook(self.workbook(self.data, 'correct'), self.expected)
        self.assertFalse(raw_result['details']); self.assertFalse(book_result['details'])
        old = json.loads((ROOT/'eval/scenarios/format-payment/expected.json').read_text())
        self.assertNotIn('Payee', [c['label'] for c in old['tables']['Payments']['columns']])

    def test_missing_swapped_and_bank_substituted_parties_fail_both_scores(self):
        for variant in ('missing', 'swapped', 'bank'):
            with self.subTest(variant=variant):
                data = copy.deepcopy(self.data); fields = data['documents'][1]['fields']
                if variant == 'missing':
                    fields.pop('payee'); fields.pop('payer')
                elif variant == 'swapped':
                    fields['payee'], fields['payer'] = fields['payer'], fields['payee']
                else:
                    fields['payer'] = fields['bank_name']
                raw_result = raw.score(data, self.expected['raw_tables'], score.score_table)
                book_result = score.score_workbook(self.workbook(data, variant), self.expected)
                for result in (raw_result, book_result):
                    self.assertLess(score.metrics(result)['field_accuracy']['rate'], 1)
                    self.assertEqual(len(result['details']), 1 if variant == 'bank' else 2)

    def test_unreadable_payer_stays_blank_and_flagged(self):
        self.data['documents'][1]['fields'].pop('payer')
        self.data['documents'][1]['unclear'] = ['payer']
        with closing(load_workbook(self.workbook(self.data, 'unclear'))) as book:
            row = score.rows(book['Payments'])[0]
            self.assertIsNone(row['Payer'])
            self.assertEqual(row['Payee'], 'Example Workshop Ltd')
            self.assertIn('Payer:', row['Review reasons'])

    def test_cash_memo_and_invoice_remain_two_billing_records(self):
        data = {'documents':[
            {'document_id':'top','file':'sale.pdf','pages':'1','document_type':'Invoice',
             'fields':{'document_title':'INVOICE','total':'12'}},
            {'document_id':'bottom','file':'sale.pdf','pages':'1','document_type':'Invoice',
             'fields':{'document_title':'CASH MEMO','total':'34'}}]}
        source = self.work/'sale.json'; source.write_text(json.dumps(data))
        output = self.work/'sale.xlsx'
        with redirect_stdout(io.StringIO()):
            self.assertEqual(process.main(['--extraction',str(source),'--out',str(output)]), 0)
        with closing(load_workbook(output)) as book:
            rows = score.rows(book['Billing documents'])
            self.assertEqual([r['Document type'] for r in rows], ['Invoice','Invoice'])
            self.assertEqual([r['Document title'] for r in rows], ['INVOICE','CASH MEMO'])
            self.assertFalse(score.rows(book['Reference documents']))

    def test_occurrence_join_handles_same_type_page_and_reordered_line_rows(self):
        book = Workbook(); docs = book.active; docs.title = 'Billing documents'
        docs.append(['File','Pages','Document ID','Document type','Total'])
        docs.append(['sale.pdf','1','top','Invoice',12])
        docs.append(['sale.pdf','1','bottom','Invoice',34])
        lines = book.create_sheet('Line items')
        lines.append(['File','Document ID','Line','Line total'])
        lines.append(['sale.pdf','bottom',1,34]); lines.append(['sale.pdf','top',1,12])
        expected = {'tables':{
            'Billing documents':{'identity':['File','Pages','Document type','Document occurrence'],
                'columns':[{'label':'Total','kind':'amount'}], 'rows':[
                    {'File':'sale.pdf','Pages':'1','Document type':'Invoice','Document occurrence':str(i),'Total':n}
                    for i,n in ((1,12),(2,34))]},
            'Line items':{'identity':['File','Document type','Document occurrence','Line'],
                'columns':[{'label':'Line total','kind':'amount'}], 'rows':[
                    {'File':'sale.pdf','Document type':'Invoice','Document occurrence':str(i),'Line':'1','Line total':n}
                    for i,n in ((1,12),(2,34))]}}}
        path = self.work/'order.xlsx'; book.save(path)
        self.assertFalse(score.score_workbook(path, expected)['details'])
        lines['D2'], lines['D3'] = 12, 34; book.save(path)
        self.assertEqual(score.score_workbook(path, expected)['cells']['silent_error'], 2)
        # An ambiguous model ID cannot provide a trustworthy occurrence join.
        docs['C3'] = 'top'; book.save(path); book.close()
        self.assertEqual(score.score_workbook(path, expected)['tables']['Line items']['missing'], 2)
