"""Read the company profile workbook and decide sales or purchase."""

import re
import tomllib
from pathlib import Path

DATA = Path(__file__).parent / "data"
SALES, PURCHASE, UNCLEAR = "Sales (AR)", "Purchase (AP)", "Unclear"
DATE_ORDER_CHOICES = ("Day/Month/Year", "Month/Day/Year", "Not sure")

_SUFFIXES = {"limited": "ltd", "company": "co", "corporation": "corp",
             "incorporated": "inc", "and": "&"}


class ProfileError(Exception):
    """A problem with the profile, explained for the person who filled it in."""


def load_columns():
    with open(DATA / "columns.toml", "rb") as f:
        return tomllib.load(f)


def normalise_name(name):
    """Compare names ignoring case, punctuation and Ltd/Limited-style variants."""
    words = re.sub(r"[.,'’()\"]", " ", str(name).casefold()).split()
    return " ".join(_SUFFIXES.get(w, w) for w in words)


def _rows(sheet):
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if any(cell not in (None, "") for cell in row):
            yield [("" if cell is None else str(cell).strip()) for cell in row]


def read_profile(path):
    """Return a dict: names (normalised set), date_order, columns (ordered export
    keys) and required (set of keys)."""
    from openpyxl import load_workbook

    path = Path(path)
    if not path.exists():
        raise ProfileError(f"The company profile '{path.name}' wasn't found.")
    try:
        book = load_workbook(path, data_only=True)
    except Exception:
        raise ProfileError(f"'{path.name}' couldn't be opened as an Excel file.")
    for needed in ("Company", "Settings", "Columns"):
        if needed not in book.sheetnames:
            raise ProfileError(
                f"'{path.name}' has no '{needed}' sheet. Start again from a fresh profile template.")

    names = set()
    for row in _rows(book["Company"]):
        for name in [row[0]] + (row[1].split(";") if len(row) > 1 else []):
            if name.strip():
                names.add(normalise_name(name))
    if not names:
        raise ProfileError("The 'Company' sheet has no company name. Add at least one in column A.")

    date_order = None
    for row in _rows(book["Settings"]):
        if row[0].lower().startswith("date order") and len(row) > 1:
            if row[1] in ("Day/Month/Year", "Month/Day/Year"):
                date_order = row[1]

    spec = {c["label"].casefold(): c for c in load_columns()["documents"]}
    columns, required = [], set()
    for row in _rows(book["Columns"]):
        column = spec.get(row[0].casefold())
        if column is None:
            raise ProfileError(
                f"The 'Columns' sheet lists '{row[0]}', which isn't a column this skill writes.")
        export = len(row) < 2 or row[1].lower() != "no"
        if export or column.get("locked"):
            columns.append(column["key"])
        if len(row) > 2 and row[2].lower() == "yes":
            required.add(column["key"])
    for column in spec.values():
        if column.get("locked") and column["key"] not in columns:
            columns.append(column["key"])

    return {"names": names, "date_order": date_order, "columns": columns,
            "required": required, "file": path.name}


def direction(profile, seller_name, buyer_name):
    """Return (direction, reason id or None)."""
    if profile is None:
        return UNCLEAR, "no_profile"
    seller = bool(seller_name) and normalise_name(seller_name) in profile["names"]
    buyer = bool(buyer_name) and normalise_name(buyer_name) in profile["names"]
    if seller and buyer:
        return UNCLEAR, "direction_both"
    if seller:
        return SALES, None
    if buyer:
        return PURCHASE, None
    return UNCLEAR, "direction_none"
