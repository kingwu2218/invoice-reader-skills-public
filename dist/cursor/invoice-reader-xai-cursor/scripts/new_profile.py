# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl>=3.1"]
# ///
"""Create a blank company profile workbook for the user to fill in.

Usage: uv run new_profile.py [path]   (default: company-profile.xlsx)
"""

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from company_profile import DATE_ORDER_CHOICES, load_columns

INSTRUCTIONS = [
    "Company profile for the invoice reader",
    "",
    "1. Company sheet: put each company you keep books for in column A.",
    "   In column B, list other names it appears under on documents, separated by ;",
    "   (for example: ABC Trading Ltd; ABC Trading Co).",
    "   If the seller on a document matches, it is a sales invoice (AR).",
    "   If the buyer matches, it is a purchase invoice (AP).",
    "2. Settings sheet: choose how dates are written on your documents.",
    "   With 'Not sure', dates such as 03/04/2026 are left for you to check.",
    "3. Columns sheet: the order of the rows is the order of the columns in the output.",
    "   Set Export to No to leave a column out. File, Needs review and Review reasons are always included.",
    "   Set Required to Yes if a document must show that field; a missing one is flagged.",
    "",
    "Save this file and tell the agent where it is.",
]


def build(path):
    book = Workbook()
    sheet = book.active
    sheet.title = "Instructions"
    for line in INSTRUCTIONS:
        sheet.append([line])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.column_dimensions["A"].width = 100

    company = book.create_sheet("Company")
    company.append(["Company name", "Other names on documents (separate with ;)"])
    company.column_dimensions["A"].width = 40
    company.column_dimensions["B"].width = 70

    settings = book.create_sheet("Settings")
    settings.append(["Setting", "Value"])
    settings.append(["Date order on your documents", "Not sure"])
    settings.column_dimensions["A"].width = 35
    settings.column_dimensions["B"].width = 20
    order = DataValidation(type="list", formula1='"' + ",".join(DATE_ORDER_CHOICES) + '"')
    settings.add_data_validation(order)
    order.add("B2")

    columns = book.create_sheet("Columns")
    columns.append(["Column", "Export", "Required"])
    yes_no = DataValidation(type="list", formula1='"Yes,No"')
    columns.add_data_validation(yes_no)
    for column in load_columns()["documents"]:
        columns.append([column["label"], "Yes", "No"])
    yes_no.add(f"B2:C{columns.max_row}")
    columns.column_dimensions["A"].width = 30

    for sheet in (company, settings, columns):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
    book.save(path)


def main():
    path = Path(sys.argv[1] if len(sys.argv) > 1 else "company-profile.xlsx")
    if path.exists():
        print(f"'{path}' already exists, so it was left as it is.")
        return 1
    build(path)
    print(f"Created '{path}'. Open it in Excel, fill in the Company sheet, and save it.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
