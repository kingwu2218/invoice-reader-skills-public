"""Write the review workbook. Numbers are real number cells; text can never run as a formula."""

from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font

# OWASP CSV Injection: a cell starting with one of these could be read as a formula.
# Full-width forms included.
_FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r", "\n", "＝", "＋", "－", "＠")
_WIDTHS = {"text": 24, "amount": 14, "number": 10, "date": 12}
_WIDE = {"review_reasons": 60, "type_evidence": 36, "description": 40,
         "seller_address": 36, "buyer_address": 36, "evidence": 36, "reason": 40}


def put(cell, value, kind="text"):
    """Write one value. Decimals and dates keep their type; everything else is safe text."""
    if value is None or value == "":
        cell.value = None
        return
    if kind in ("amount", "number") and isinstance(value, Decimal):
        cell.value = value
        cell.number_format = "#,##0.00" if kind == "amount" else "General"
        return
    if kind == "date" and isinstance(value, date):
        cell.value = value
        cell.number_format = "yyyy-mm-dd"
        return
    text = ILLEGAL_CHARACTERS_RE.sub("", str(value))
    cell.value = text
    cell.data_type = "s"
    if text.startswith(_FORMULA_STARTS):
        # Excel's own marker for "typed with a leading apostrophe": shown and kept as text.
        cell.quotePrefix = True


def _sheet(book, title, columns, rows):
    """columns: list of (key, label, kind). rows: list of dicts keyed by column key."""
    sheet = book.create_sheet(title)
    for index, (key, label, kind) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index)
        put(cell, label)
        cell.font = Font(bold=True)
        letter = cell.column_letter
        sheet.column_dimensions[letter].width = _WIDE.get(key, _WIDTHS.get(kind, 20))
    for row_number, row in enumerate(rows, start=2):
        for index, (key, _label, kind) in enumerate(columns, start=1):
            cell = sheet.cell(row=row_number, column=index)
            put(cell, row.get(key), kind)
            if key == "review_reasons" or key.endswith("address"):
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"
    return sheet


def write(path, sheets, run_info):
    """sheets: list of (title, columns, rows), written in order, then Run info."""
    book = Workbook()
    book.remove(book.active)
    for title, columns, rows in sheets:
        _sheet(book, title, columns, rows)
    info = book.create_sheet("Run info")
    for label, value in run_info:
        info.append([None, None])
        put(info.cell(row=info.max_row, column=1), label)
        put(info.cell(row=info.max_row, column=2), value)
    info.column_dimensions["A"].width = 36
    info.column_dimensions["B"].width = 60
    book.save(path)
